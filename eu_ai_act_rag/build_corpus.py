#!/usr/bin/env python3
"""RAG-Dokumentenbestand zur KI-Verordnung (EU) 2024/1689 aus amtlichen EUR-Lex-Texten erzeugen.

Aufruf (Beispiele):
    python build_corpus.py                        # alles abrufen, erzeugen, prüfen
    python build_corpus.py --html-dir ./html      # manuell gespeicherte EUR-Lex-HTML-Dateien verwenden
    python build_corpus.py --offline              # nur Cache verwenden
    python build_corpus.py --entwuerfe            # zusätzlich Leitlinien-Entwürfe (gekennzeichnet)
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))

import config  # noqa: E402
import leitlinien  # noqa: E402
import planner  # noqa: E402
import qa  # noqa: E402
from docx_writer import write_docx  # noqa: E402
from eurlex_parser import parse  # noqa: E402
from fetch import FetchError, fetch  # noqa: E402

COLUMNS = ["Dateiname", "Ordner", "Dokumenttyp", "Kapitel/Abschnitt", "Artikel von", "Artikel bis",
           "Erwägungsgründe von", "Erwägungsgründe bis", "Anhang", "CELEX-Quelle", "Fassungsstand",
           "Seiten (geschätzt)", "Seiten (Word/PDF)", "Zeichen", "Status der Qualitätsprüfung", "Prüfhinweise"]


def row(d: planner.PlannedDoc, pages: dict) -> list:
    kap = " / ".join(x for x in (d.kapitel, d.abschnitt) if x)
    return [d.spec.path.name, d.ordner, d.dokumenttyp, kap,
            d.artikel[0] if d.artikel else "", d.artikel[-1] if d.artikel else "",
            d.erwaegungsgruende[0] if d.erwaegungsgruende else "", d.erwaegungsgruende[-1] if d.erwaegungsgruende else "",
            d.anhang, d.celex, d.fassung, round(d.zeichen / config.ZEICHEN_PRO_SEITE, 1) if d.zeichen else "",
            pages.get(d.spec.path, ""), d.zeichen, d.qa_status, "; ".join(d.qa_details)]


def write_overview(out: Path, docs: list[planner.PlannedDoc], pages: dict):
    rows = [row(d, pages) for d in docs]
    with open(out / "00_Dokumentenuebersicht.csv", "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(COLUMNS)
        w.writerows(rows)
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Dokumente"
    ws.append(COLUMNS)
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="top")
    widths = [62, 34, 28, 60, 10, 10, 12, 12, 9, 22, 40, 10, 10, 10, 14, 60]
    for i, wdt in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = wdt
    status_col = COLUMNS.index("Status der Qualitätsprüfung") + 1
    for r in range(2, ws.max_row + 1):
        c = ws.cell(r, status_col)
        c.fill = PatternFill("solid", fgColor="C6EFCE" if c.value == "OK" else "FFC7CE")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(out / "00_Dokumentenuebersicht.xlsx")


def write_report(out: Path, res: qa.QAResult, sources: dict, gl_results, stamp: str):
    lines = [f"# Qualitätsbericht RAG-Bestand KI-Verordnung", "", f"Erstellt: {stamp}", "", "## Quellen", ""]
    for fd in sources.values():
        lines.append(f"- {fd.source.beschreibung} – CELEX {fd.source.celex} – {fd.source.page_url()} – "
                     f"abgerufen {fd.abgerufen_am} ({fd.herkunft}), SHA-256 {fd.sha256}")
    lines += ["", f"## Ergebnis: {len(res.fehler)} Fehler, {len(res.warnungen)} Warnungen, "
                  f"{sum(c.ok for c in res.checks)} bestandene Prüfungen", ""]
    lines += ["| Bereich | Prüfung | Ergebnis | Details |", "|---|---|---|---|"]
    for c in res.checks:
        erg = "OK" if c.ok else c.schwere
        det = c.details.replace("|", "\\|").replace("\n", " ")[:400]
        lines.append(f"| {c.bereich} | {c.pruefung} | {erg} | {det} |")
    if gl_results:
        lines += ["", "## Leitlinien der Kommission", ""]
        for g in gl_results:
            lines.append(f"- {'OK' if g.ok else 'FEHLER'} – {g.eintrag['titel']} ({g.eintrag['status']}): {g.meldung}")
    (out / "00_Qualitaetsbericht.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    (out / "00_Qualitaetsbericht.json").write_text(json.dumps(
        [c.__dict__ for c in res.checks], ensure_ascii=False, indent=1), encoding="utf-8")


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", type=Path, default=HERE / "output" / "EU_KI_Verordnung_RAG")
    ap.add_argument("--cache", type=Path, default=HERE / "cache")
    ap.add_argument("--html-dir", type=Path, default=None,
                    help="Ordner mit manuell von EUR-Lex gespeicherten Dateien <CELEX>.html")
    ap.add_argument("--offline", action="store_true")
    ap.add_argument("--refresh", action="store_true", help="Cache ignorieren und neu abrufen")
    ap.add_argument("--ohne-leitlinien", action="store_true")
    ap.add_argument("--entwuerfe", action="store_true", help="Leitlinien-Entwürfe (gekennzeichnet) aufnehmen")
    ap.add_argument("--leitlinien-discovery", action="store_true",
                    help="Kommissionsseiten nach weiteren Leitlinien durchsuchen (nur Liste zur Prüfung)")
    ap.add_argument("--seitenzahl", action="store_true", help="echte Seitenzahl per LibreOffice ermitteln")
    ap.add_argument("--leitlinien-pdf-dir", type=Path, default=None,
                    help="Ordner mit manuell von der Kommissionsseite geladenen PDFs <id>.pdf (ids siehe leitlinien_quellen.json)")
    ap.add_argument("--quellen", type=Path, default=HERE / "leitlinien_quellen.json")
    args = ap.parse_args(argv)

    out: Path = args.out.resolve()
    if out.exists():
        shutil.rmtree(out)
    out.mkdir(parents=True)
    stamp = dt.datetime.now().isoformat(timespec="seconds")

    sources = {}
    for key, src in config.SOURCES.items():
        try:
            sources[key] = fetch(src, args.cache / "eurlex", args.html_dir, args.offline, args.refresh)
            print(f"[Quelle] {src.celex}: {sources[key].herkunft}, {len(sources[key].html):,} Zeichen HTML")
        except FetchError as exc:
            print(f"[FEHLER] {exc}", file=sys.stderr)
            return 2

    acts = {k: parse(fd.html, track_quotes=(k == "omnibus")) for k, fd in sources.items()}
    kons, orig, omni = acts["konsolidiert"], acts["original"], acts["omnibus"]

    docs_kons = planner.plan_articles(sources["konsolidiert"], kons, out)
    docs_anx = planner.plan_annexes(sources["konsolidiert"], kons, out)
    docs_rct = planner.plan_recitals(sources["original"], orig, out, "VO_2024-1689_Originalfassung", "",
                                     "Verordnung (EU) 2024/1689")
    docs_omni_rct = planner.plan_recitals(sources["omnibus"], omni, out, "VO_2026-1744_Digital-Omnibus",
                                          "VO_2026-1744_", "Verordnung (EU) 2026/1744")
    docs_omni = planner.plan_amending_act(sources["omnibus"], omni, out, "VO_2026-1744_", "Verordnung (EU) 2026/1744")
    all_docs = docs_kons + docs_anx + docs_rct + docs_omni_rct + docs_omni

    for d in all_docs:
        write_docx(d.spec)

    res = qa.QAResult()
    qa.check_consolidated(res, sources["konsolidiert"].html, kons, docs_kons + docs_anx)
    qa.check_recitals(res, "VO 2024/1689 (Original) – Erwägungsgründe", sources["original"].html, orig, docs_rct,
                      config.EXPECTED["original"]["erwaegungsgruende"])
    omni_required = [b for a in omni.articles for b in a.blocks] + omni.final + [b for a in omni.annexes for b in a.blocks]
    qa.check_recitals(res, "VO 2026/1744 – Erwägungsgründe und verfügender Teil", sources["omnibus"].html, omni,
                      docs_omni_rct + docs_omni, config.EXPECTED["omnibus"]["erwaegungsgruende"], omni_required)
    qa.check_articles_vs_eli(res, "VO 2026/1744 – Erwägungsgründe und verfügender Teil", sources["omnibus"].html, omni)
    for d in all_docs:
        qa.check_docx(res, d)

    gl_results = []
    if not args.ohne_leitlinien:
        gl_results = leitlinien.build(out, args.cache / "leitlinien", args.entwuerfe, args.quellen,
                                       args.leitlinien_pdf_dir)
        for g in gl_results:
            res.add("Leitlinien", g.eintrag["titel"], g.ok, g.meldung, "WARNUNG")
            all_docs += g.docs
    if args.leitlinien_discovery:
        cand = leitlinien.discover(args.cache / "leitlinien", args.quellen)
        with open(out / "00_Leitlinien_Kandidaten_zur_Pruefung.csv", "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(["Linktext", "URL", "Status (heuristisch)"])
            w.writerows(cand)

    # Dokumentstatus aus globalen Fehlern der jeweiligen Quelle ableiten
    global_fail = {c.bereich for c in res.fehler}
    for d in all_docs:
        src_area = {"konsolidiert": "Konsolidierte Fassung", "original": "VO 2024/1689 (Original) – Erwägungsgründe",
                    "omnibus": "VO 2026/1744 – Erwägungsgründe und verfügender Teil"}.get(d.quelle_key)
        if src_area in global_fail and d.qa_status == "OK":
            d.qa_status = "WARNUNG"
            d.qa_details.append(f"Globale Prüfung der Quelle nicht bestanden (siehe Qualitätsbericht: {src_area})")

    pages = qa.count_pages([d.spec.path for d in all_docs]) if args.seitenzahl else {}
    write_overview(out, all_docs, pages)
    write_report(out, res, sources, gl_results, stamp)

    print(f"\nAusgabeordner: {out}")
    for d in all_docs:
        print(f"  [{d.qa_status:7}] {d.spec.path.relative_to(out)}")
    print(f"\nQualitätsprüfung: {len(res.fehler)} Fehler, {len(res.warnungen)} Warnungen – Details in 00_Qualitaetsbericht.md")
    for c in res.fehler[:30]:
        print(f"  FEHLER {c.bereich}: {c.pruefung} {c.details[:200]}")
    return 1 if res.fehler else 0


if __name__ == "__main__":
    sys.exit(main())
