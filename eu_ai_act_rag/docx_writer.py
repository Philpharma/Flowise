"""Erzeugung der Word-Dateien aus den geparsten Blöcken (ohne jede inhaltliche Veränderung).

Konventionen, auf die sich auch die Qualitätsprüfung stützt:
* Absätze mit Formatvorlagen, deren Name mit "RAG " beginnt, sind technische Zusätze
  (Metadaten, wiederholte Kontextüberschriften bei Teildokumenten, Fußnoten-Überschrift) und
  gehören nicht zum amtlichen Text.
* Die Zeichenformatvorlage "RAG Tabellenlabel" kennzeichnet Spaltenbezeichnungen, die bei der
  Umwandlung komplexer Tabellen in strukturierten Text zur Zuordnung wiederholt werden.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from eurlex_parser import Block, Run

META_STYLE = "RAG Metadaten"
CONTEXT_STYLES = {1: "RAG Kontext 1", 2: "RAG Kontext 2", 3: "RAG Kontext 3"}
NOTE_HEAD_STYLE = "RAG Fussnoten Ueberschrift"
NOTE_STYLE = "RAG Fussnote"
DRAFT_STYLE = "RAG Entwurfshinweis"
LABEL_STYLE = "RAG Tabellenlabel"
INDENT_CM = 0.75


@dataclass
class Section:
    """Ein Element des Dokumentinhalts in Reihenfolge."""
    kind: str  # "heading" | "blocks" | "context"
    level: int = 1
    blocks: list[Block] = field(default_factory=list)
    text: str = ""  # nur für "context"


@dataclass
class DocSpec:
    path: Path
    titel: str
    metadata: list[tuple[str, str]]
    sections: list[Section]
    notes: list[Block] = field(default_factory=list)
    entwurf: bool = False
    core_subject: str = ""
    core_keywords: str = ""

    def content_blocks(self) -> list[Block]:
        return [b for s in self.sections if s.kind in ("heading", "blocks") for b in s.blocks]


def _ensure_styles(doc: Document):
    st = doc.styles
    normal = st["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    pf = normal.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(4), 1.05
    for lvl, size in ((1, 15), (2, 13), (3, 11.5)):
        h = st[f"Heading {lvl}"]
        h.font.size = Pt(size)
        h.font.name = "Calibri"
        h.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
        h.paragraph_format.space_before = Pt(12 if lvl < 3 else 9)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
    if META_STYLE not in [s.name for s in st]:
        m = st.add_style(META_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        m.base_style = normal
        m.font.size = Pt(8.5)
        m.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        m.paragraph_format.space_after = Pt(1)
        for lvl, name in CONTEXT_STYLES.items():
            c = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
            c.base_style = st[f"Heading {lvl}"]
            c.font.italic = True
            c.font.color.rgb = RGBColor(0x6B, 0x7B, 0x8C)
        nh = st.add_style(NOTE_HEAD_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        nh.base_style = st["Heading 3"]
        n = st.add_style(NOTE_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        n.base_style = normal
        n.font.size = Pt(9)
        d = st.add_style(DRAFT_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        d.base_style = normal
        d.font.bold = True
        d.font.size = Pt(14)
        d.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
        lab = st.add_style(LABEL_STYLE, WD_STYLE_TYPE.CHARACTER)
        lab.font.bold = True
        lab.font.color.rgb = RGBColor(0x44, 0x44, 0x44)


def _page_setup(doc: Document):
    for s in doc.sections:
        s.page_height, s.page_width = Cm(29.7), Cm(21.0)
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)


def _add_runs(par, runs: list[Run]):
    for r in runs:
        run = par.add_run(r.text)
        if r.sup:
            run.font.superscript = True
        if r.sub:
            run.font.subscript = True
        if r.bold:
            run.bold = True
        if r.italic:
            run.italic = True


def _write_para(doc: Document, b: Block, style: str | None = None, base: int = 0):
    par = doc.add_paragraph(style=style)
    pf = par.paragraph_format
    indent = Cm(INDENT_CM * max(b.level - base, 0))
    if b.enum:
        hang = Cm(INDENT_CM if len(b.enum) <= 4 else 1.1)
        pf.left_indent = indent + hang
        pf.first_line_indent = -hang
        pf.tab_stops.add_tab_stop(indent + hang)
        par.add_run(b.enum)
        par.add_run("\t")
    elif b.level:
        pf.left_indent = indent
    _add_runs(par, b.runs)
    return par


def _heading(doc: Document, blocks: list[Block], level: int, style: str | None = None):
    par = doc.add_paragraph(style=style or f"Heading {level}")
    for i, b in enumerate(blocks):
        if i:
            par.add_run().add_break(WD_BREAK.LINE)
        if b.enum:
            par.add_run(b.enum + " ")
        _add_runs(par, b.runs)
    return par


def _set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def _write_table(doc: Document, b: Block):
    rows = b.rows
    if not b.is_complex_table:
        ncols = max(len(r) for r in rows)
        t = doc.add_table(rows=len(rows), cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, r in enumerate(rows):
            for j, c in enumerate(r):
                cell = t.cell(i, j)
                cell.text = ""
                lines = c.text.split("\n")
                p = cell.paragraphs[0]
                for k, line in enumerate(lines):
                    if k:
                        p = cell.add_paragraph()
                    run = p.add_run(line)
                    run.bold = c.header or (i == 0 and all(x.header for x in r))
                    run.font.size = Pt(9.5)
        if rows and all(c.header for c in rows[0]):
            _set_repeat_header(t.rows[0])
        doc.add_paragraph()
        return
    # Komplexe Tabelle (verbundene Zellen): strukturierter Text, Zeile für Zeile, ohne Informationsverlust.
    doc.add_paragraph("Tabelle (in strukturierten Text überführt, verbundene Zellen je Zeile aufgelöst):", style=META_STYLE)
    grid: list[list[tuple[str, bool]]] = []  # (Text, ist_Wiederholung)
    spans: dict[tuple[int, int], tuple[str, int]] = {}
    for i, r in enumerate(rows):
        line: list[tuple[str, bool]] = []
        col = 0
        it = iter(r)
        while True:
            while (i, col) in spans:
                txt, left = spans.pop((i, col))
                line.append((txt, True))
                if left > 1:
                    spans[(i + 1, col)] = (txt, left - 1)
                col += 1
            c = next(it, None)
            if c is None:
                break
            for k in range(c.colspan):
                line.append((c.text, k > 0))
                if c.rowspan > 1:
                    spans[(i + 1, col)] = (c.text, c.rowspan - 1)
                col += 1
        grid.append(line)
    header_rows = 0
    for r in rows:
        if all(c.header or not c.text for c in r) and any(c.header for c in r):
            header_rows += 1
        else:
            break
    if header_rows == 0 and len(rows) > 1:
        header_rows = 1
    ncols = max(len(line) for line in grid)
    headers = []
    for j in range(ncols):
        parts = []
        for line in grid[:header_rows]:
            if j < len(line) and line[j][0] and line[j][0] not in parts:
                parts.append(line[j][0])
        headers.append(" / ".join(parts))
    for line in grid[:header_rows]:
        p = doc.add_paragraph()
        p.add_run("Tabellenkopf: ", style=LABEL_STYLE)
        for j, (txt, rep) in enumerate(line):
            if j:
                p.add_run(" | ", style=LABEL_STYLE)
            p.add_run(txt, style=LABEL_STYLE if rep else None)
    for i, line in enumerate(grid[header_rows:], start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(INDENT_CM)
        p.add_run(f"Zeile {i}: ", style=LABEL_STYLE)
        for j, (txt, rep) in enumerate(line):
            if j:
                p.add_run(" | ", style=LABEL_STYLE)
            label = headers[j] if j < len(headers) and headers[j] else f"Spalte {j + 1}"
            p.add_run(f"{label}: ", style=LABEL_STYLE)
            p.add_run(txt, style=LABEL_STYLE if rep else None)


def write_docx(spec: DocSpec) -> Path:
    doc = Document()
    _ensure_styles(doc)
    _page_setup(doc)
    cp = doc.core_properties
    cp.title = spec.titel[:250]
    cp.subject = spec.core_subject[:250]
    cp.keywords = spec.core_keywords[:250]
    cp.author = "Automatisch aus EUR-Lex erzeugt (Python)"
    cp.comments = "; ".join(f"{k}: {v}" for k, v in spec.metadata[1:3])[:250]

    if spec.entwurf:
        doc.add_paragraph("ENTWURF / DRAFT – keine endgültig angenommene Fassung", style=DRAFT_STYLE)
    doc.add_paragraph(f"Dokument: {spec.titel}", style=META_STYLE)
    for k, v in spec.metadata:
        doc.add_paragraph(f"{k}: {v}", style=META_STYLE)

    base = min((b.level for b in spec.content_blocks() if b.kind == "p"), default=0)
    for s in spec.sections:
        if s.kind == "context":
            doc.add_paragraph(s.text, style=CONTEXT_STYLES.get(s.level, CONTEXT_STYLES[3]))
        elif s.kind == "heading":
            _heading(doc, s.blocks, s.level)
        else:
            for b in s.blocks:
                if b.kind == "table":
                    _write_table(doc, b)
                else:
                    _write_para(doc, b, base=base)
    if spec.notes:
        doc.add_paragraph("Fußnoten", style=NOTE_HEAD_STYLE)
        for n in spec.notes:
            _write_para(doc, n, style=NOTE_STYLE)
    zoom = doc.settings.element.find(qn("w:zoom"))
    if zoom is not None and zoom.get(qn("w:percent")) is None:
        zoom.set(qn("w:percent"), "100")  # Schemakonformität der python-docx-Vorlage
    spec.path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(spec.path)
    return spec.path


# ------------------------------------------------------------------------------------------------
# Rücklesen für die Qualitätsprüfung
# ------------------------------------------------------------------------------------------------
def read_back(path: Path) -> dict:
    """Liest eine erzeugte Datei und trennt amtlichen Inhalt, Fußnoten und technische Zusätze."""
    doc = Document(path)
    body = doc.element.body
    content, notes, meta = [], [], []
    for child in body.iterchildren():
        tag = child.tag.split("}")[1]
        if tag == "p":
            from docx.text.paragraph import Paragraph
            par = Paragraph(child, doc)
            style = par.style.name if par.style is not None else ""
            text = "".join(r.text for r in _iter_runs(par) if not (r.style is not None and r.style.name == LABEL_STYLE))
            if style == NOTE_STYLE:
                notes.append(text)
            elif style.startswith("RAG "):
                meta.append(par.text)
            else:
                content.append(text)
        elif tag == "tbl":
            from docx.table import Table
            tbl = Table(child, doc)
            for row in tbl.rows:
                seen = set()
                for cell in row.cells:
                    if id(cell._tc) in seen:
                        continue
                    seen.add(id(cell._tc))
                    content.append("\n".join(p.text for p in cell.paragraphs))
    return {"content": content, "notes": notes, "meta": meta}


def _iter_runs(par):
    # Absatz-Runs inkl. Zeilenumbrüchen/Tabs; Text der Runs genügt für den Vergleich.
    return par.runs
