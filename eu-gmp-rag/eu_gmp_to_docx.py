#!/usr/bin/env python3
"""
EU-GMP-Leitfaden (BMG-Bekanntmachungen) -> RAG-taugliche Word-Dateien.

Ablauf:
  1. Seite https://www.bundesgesundheitsministerium.de/service/gesetze-und-verordnungen/bekanntmachungen
     laden und alle PDF-Links des Abschnitts "EU-GMP Leitfaden" einsammeln
     (oder mit --pdf-dir bereits heruntergeladene PDFs verwenden).
  2. Text wortgetreu extrahieren, dabei entfernen:
       - Kopf-/Fußzeilen und Seitenzahlen (Zeilen, die sich auf vielen Seiten wiederholen)
       - die Dokumenthistorie-Tabelle am Anfang
       - das Deckblatt der Bekanntmachung (Text vor der ersten nummerierten Überschrift)
  3. Tabellen im Fließtext zeilenweise in Sätze umwandeln
     ("Spalte A: Wert; Spalte B: Wert"), damit ein RAG sie lesen kann.
  4. Lange Dokumente (Teil II, Teil IV, große Anhänge) an Kapitelgrenzen in
     Teile von max. ca. 15 Seiten aufteilen und sprechend benennen.
  5. Jede Word-Datei beginnt mit Titel + Quellenzeile; Überschriften sind echte
     Word-Überschriften (Überschrift 1/2), damit Chunker sie erkennen.

Aufruf:
  python eu_gmp_to_docx.py --out ausgabe/            # lädt alles von der BMG-Seite
  python eu_gmp_to_docx.py --pdf-dir pdfs/ --out ausgabe/   # nutzt lokale PDFs
"""
import argparse
import re
import statistics
import sys
import zipfile
from collections import Counter
from pathlib import Path
from urllib.parse import urljoin

import pdfplumber
from docx import Document
from docx.shared import Pt

SOURCE_PAGE = "https://www.bundesgesundheitsministerium.de/service/gesetze-und-verordnungen/bekanntmachungen"

# ~500 Wörter je Word-Seite (A4, 11 pt) -> 6.500 Wörter ≈ 13 Seiten, Luft bis 15 Seiten.
MAX_WORDS = 6500

# Sprechende Namen (Fallback, falls der Titel im PDF nicht sauber erkennbar ist)
TITLES = {
    "Einleitung": "Einleitung EU-GMP-Leitfaden",
    "Kapitel 1": "Teil I Kapitel 1 Pharmazeutisches Qualitätssystem",
    "Kapitel 2": "Teil I Kapitel 2 Personal",
    "Kapitel 3": "Teil I Kapitel 3 Räumlichkeiten und Ausrüstung",
    "Kapitel 4": "Teil I Kapitel 4 Dokumentation",
    "Kapitel 5": "Teil I Kapitel 5 Produktion",
    "Kapitel 6": "Teil I Kapitel 6 Qualitätskontrolle",
    "Kapitel 7": "Teil I Kapitel 7 Ausgelagerte Tätigkeiten",
    "Kapitel 8": "Teil I Kapitel 8 Beanstandungen, Qualitätsmängel und Produktrückrufe",
    "Kapitel 9": "Teil I Kapitel 9 Selbstinspektionen",
    "Teil II": "Teil II Wirkstoffe",
    "Anhang 3": "Anhang 3 Herstellung von Radiopharmaka",
    "Anhang 6": "Anhang 6 Herstellung medizinischer Gase",
    "Anhang 7": "Anhang 7 Herstellung pflanzlicher Arzneimittel",
    "Anhang 8": "Anhang 8 Probenahme von Ausgangs- und Verpackungsmaterial",
    "Anhang 9": "Anhang 9 Herstellung von Liquida, Cremes und Salben",
    "Anhang 10": "Anhang 10 Herstellung von Dosieraerosolen zur Inhalation",
    "Anhang 11": "Anhang 11 Computergestützte Systeme",
    "Anhang 12": "Anhang 12 Ionisierende Strahlung in der Arzneimittelherstellung",
    "Anhang 13": "Anhang 13 Prüfpräparate",
    "Anhang 14": "Anhang 14 Arzneimittel aus menschlichem Blut oder Plasma",
    "Anhang 15": "Anhang 15 Qualifizierung und Validierung",
    "Anhang 16": "Anhang 16 Zertifizierung durch die Sachkundige Person und Chargenfreigabe",
    "Anhang 19": "Anhang 19 Referenz- und Rückstellmuster",
    "Teil IV": "Teil IV ATMP",
}

LINK_RE = re.compile(r"^(Einleitung|Kapitel \d+|Teil [IV]+|Anhang \d+)\b")
HISTORY_KEYWORDS = ("historie", "history", "geschichte des dokuments", "grund der änderung", "überarbeitung",
                    "inkrafttreten", "datum der", "revision", "stand der")
PAGE_NO_RE = re.compile(r"^(seite\s*)?\d{1,3}(\s*(von|/)\s*\d{1,3})?$", re.I)
H1_RE = re.compile(r"^(\d{1,2})\.?\s+([A-ZÄÖÜ][^.]{2,120})$")
H2_RE = re.compile(r"^(\d{1,2})\.(\d{1,2})\.?\s+([A-ZÄÖÜ][^.]{2,120})$")


# --------------------------------------------------------------------------- Download
def download_pdfs(target: Path) -> list[tuple[str, Path]]:
    import requests
    from bs4 import BeautifulSoup

    target.mkdir(parents=True, exist_ok=True)
    html = requests.get(SOURCE_PAGE, timeout=60).text
    soup = BeautifulSoup(html, "html.parser")

    # Abschnitt "EU-GMP Leitfaden" suchen (Akkordeon-Überschrift) und seine Links nehmen
    anchor = soup.find(string=re.compile(r"EU-GMP[- ]Leitfaden"))
    if anchor is None:
        sys.exit("Abschnitt 'EU-GMP Leitfaden' auf der Seite nicht gefunden.")
    container = anchor.find_parent()
    while container is not None and len([a for a in container.find_all("a") if ".pdf" in (a.get("href") or "")]) < 5:
        container = container.find_parent()

    result, seen = [], set()
    for a in container.find_all("a"):
        href = a.get("href") or ""
        text = " ".join(a.get_text(" ").split())
        m = LINK_RE.match(text)
        if not m or ".pdf" not in href.lower() or m.group(1) in seen:
            continue
        seen.add(m.group(1))
        url = urljoin(SOURCE_PAGE, href)
        path = target / f"{m.group(1)}.pdf"
        if not path.exists():
            r = requests.get(url, timeout=120)
            r.raise_for_status()
            path.write_bytes(r.content)
        print(f"  geladen: {m.group(1):12s} <- {url}")
        result.append((m.group(1), path))
    return result


# --------------------------------------------------------------------------- Extraktion
def clean(text: str) -> str:
    """Fehlkodierte Sonderzeichen (Bundesanzeiger-PDFs) reparieren."""
    text = re.sub(r"\(cid:2[24]\)\(cid:6\)\s*", "≥ ", text)
    text = re.sub(r"\(cid:23\)\(cid:6\)\s*", "< ", text)
    text = re.sub("[\uf000-\uf0ff\u25aa]", "•", text)  # Symbol-/Wingdings-Aufzählungszeichen
    return re.sub(r"\(cid:\d+\)", "", text)


def table_to_sentences(rows: list[list[str | None]]) -> list[str]:
    """Tabelle -> eine Zeile Text pro Tabellenzeile, mit Spaltenköpfen als Schlüssel."""
    rows = [[" ".join(clean(c or "").replace("-\n", "").split()) for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    # Zellen, deren Text über mehrere Tabellenzeilen verteilt ist, wieder zusammenführen:
    # Zeilen ohne Eintrag in der ersten Spalte sind Fortsetzungen der Zeile darüber
    merged = []
    for r in rows:
        if merged and not r[0]:
            prev = merged[-1]
            for j, c in enumerate(r):
                if c:
                    p_ = prev[j].replace(GREY, "")
                    joined = p_[:-1] + c if p_.endswith("-") and c[:1].islower() else f"{p_} {c}".strip()
                    prev[j] = joined + (GREY if GREY in prev[j] + c else "")
        else:
            merged.append(r)
    rows = [[(c.replace(GREY, "") + (" (grau markiert)" if GREY in c else "")).strip() for c in r] for r in merged]
    if not rows:
        return []
    # horizontal verbundene Zellen im Kopf: Text nach rechts übernehmen
    header = rows[0][:]
    for i in range(1, len(header)):
        if not header[i]:
            header[i] = header[i - 1]
    body = rows[1:]
    if not body or not any(header):
        return [" | ".join(c for c in r if c) for r in rows]
    out = []
    for r in body:
        parts, last = [], None
        for h, c in zip(header, r):
            if not c:
                continue
            if h and h == last:
                if not parts[-1].endswith(c):
                    parts[-1] += f", {c}"
            else:
                parts.append(f"{h}: {c}" if h else c)
            last = h
        if parts:
            out.append("; ".join(parts) + ("" if parts[-1].endswith(".") else "."))
    return out


GREY = "\x00"  # Markierung für grau hinterlegte Zellen


def is_grey(color) -> bool:
    if not color or not isinstance(color, (tuple, list)):
        return False
    vals = list(color) if len(color) in (1, 3) else []
    return bool(vals) and max(vals) - min(vals) < 0.05 and 0.3 < vals[0] < 0.97


def extract_table(table, page) -> list[list[str | None]]:
    """Tabelle auslesen; grau hinterlegte Zellen markieren (z. B. Teil II, Tabelle 1: 'grau markiert' = GMP gilt)."""
    rows = table.extract(x_tolerance=1.5)
    greys = [r for r in page.rects if r.get("fill") and is_grey(r.get("non_stroking_color"))]
    for row, robj in zip(rows, table.rows):
        for j, (c, box) in enumerate(zip(row, robj.cells)):
            if c is not None and box:
                cx, cy = (box[0] + box[2]) / 2, (box[1] + box[3]) / 2
                if any(g["x0"] <= cx <= g["x1"] and g["top"] <= cy <= g["bottom"] for g in greys):
                    row[j] = (c or "") + GREY
    return rows


def is_history_table(rows) -> bool:
    text = " ".join(" ".join((c or "").replace(GREY, "") for c in r) for r in rows[:2]).lower()
    return any(k in text for k in HISTORY_KEYWORDS)


BOLD_RE = re.compile(r"bold|black|heavy|[-.,]bd?$|[-.,]bd?[-,]")
NOISE_RE = re.compile(r"^(Bekanntmachung|Veröffentlicht am .*|BAnz AT [\d.]+ B\d+|www\.bundesanzeiger\.de.*|-\s*\d+\s*-)$"
                      r"|Westferry Circus|Churchill Place|Domenico Scarlattilaan|^Telephone \+|Send (us )?a question|"
                      r"An agency of the European Union|© European Medicines Agency|^Official address|^Address for visits|"
                      r"^E -?mail info@ema")


def extract_blocks(pdf_path: Path, line_numbers: bool = False) -> list[tuple[str, object]]:
    """Liefert eine Liste von (typ, payload); typ in {'line', 'table', 'gap'}.
    line_numbers=True entfernt Zeilennummern am Zeilenanfang (z. B. Kommissions-Leitlinie Prüfpräparate)."""
    pages = []
    with pdfplumber.open(pdf_path) as pdf:
        for pno, page in enumerate(pdf.pages):
            items = []  # (top, typ, payload)
            # echte Tabellen: mind. 2 Spalten und 2 Zeilen; einspaltige Rahmen/Kästen sind Fließtext
            found = [(t.bbox, extract_table(t, page)) for t in page.find_tables()]
            found = [(b, r) for b, r in found if len(r) >= 2 and max(len(x) for x in r) >= 2]
            tables = [(b, r) for b, r in found
                      if not any(o != b and o[0] <= b[0] and o[1] <= b[1] and b[2] <= o[2] and b[3] <= o[3]
                                 for o, _ in found)]
            # Tabellen, deren Rahmen nicht alle Spalten umfasst (Text links daneben, z. B. Reinraumklassen
            # in Teil IV), zeilenweise als Text übernehmen statt sie zu zerreißen
            def partial(b):
                return any(c["x1"] < b[0] - 2 and b[1] <= (c["top"] + c["bottom"]) / 2 <= b[3] and c["text"].strip()
                           for c in page.chars)
            bands = [(b[1], b[3]) for b, _ in tables if partial(b)]
            tables = [(b, r) for b, r in tables if not partial(b)]
            bboxes = [b for b, _ in tables]
            for b, rows in tables:
                if pno < 2 and is_history_table(rows):
                    continue  # Dokumenthistorie am Anfang verwerfen
                items.append((b[1], "table", rows))

            def outside(obj):
                return not any(b[0] - 1 <= obj["x0"] and obj["x1"] <= b[2] + 1 and
                               b[1] - 1 <= obj["top"] and obj["bottom"] <= b[3] + 1 for b in bboxes)

            filtered = page.filter(lambda o: o.get("object_type") != "char" or outside(o))
            lines = filtered.extract_text_lines(layout=False, strip=True, return_chars=True, x_tolerance=1.5)
            sizes = [c["size"] for ln in lines for c in ln["chars"] if c["text"].strip()]
            body_size = statistics.median(sizes) if sizes else 10
            # Spaltenköpfe oberhalb solcher Tabellen mitnehmen (bis zur einleitenden Zeile "…aufgeführt:")
            for k, (t0, t1) in enumerate(bands):
                for ln in sorted((l for l in lines if l["bottom"] <= t0 + 1 and l["top"] > t0 - 50), key=lambda l: -l["top"]):
                    if ln["text"].rstrip().endswith(":") or PARA_NO_RE.match(ln["text"]):
                        break
                    t0 = ln["top"]
                bands[k] = (t0, t1)
            for ln in lines:
                if ln["text"].count("(cid:") >= 8:
                    continue  # unlesbare Verlagszeile
                chars = [c for c in ln["chars"] if c["text"].strip()]
                if not chars:
                    continue
                bold = sum(bool(BOLD_RE.search(c["fontname"].lower())) for c in chars) > len(chars) / 2
                size = statistics.median(c["size"] for c in chars)
                # Fußnote: kleiner als der Fließtext und in der unteren Seitenhälfte
                foot = size < body_size * 0.88 and ln["top"] > page.height * 0.5
                tab = any(t0 - 1 <= ln["top"] <= t1 + 1 for t0, t1 in bands)
                text = clean(ln["text"]).strip()
                if line_numbers and not foot:
                    text = re.sub(r"^\d{1,4}\s+", "", text)
                items.append((ln["top"], "foot" if foot else ("tabline" if tab else "line"),
                              (text, ln["bottom"], page.height, bold)))
            items.sort(key=lambda x: x[0])
            pages.append(items)

    # Kopf-/Fußzeilen: Zeilen im oberen/unteren 10 % der Seite, die auf >= 40 % der Seiten vorkommen
    norm = lambda s: re.sub(r"\d+", "#", s.strip().lower())
    counter = Counter()
    for items in pages:
        seen = set()
        for top, typ, p in items:
            if typ != "table" and (top < p[2] * 0.1 or top > p[2] * 0.9):
                seen.add(norm(p[0]))
        counter.update(seen)
    repeated = {k for k, v in counter.items() if len(pages) > 2 and v >= max(2, 0.4 * len(pages))}

    blocks = []
    for items in pages:
        heights = [p[1] - top for top, typ, p in items if typ != "table"]
        lh = statistics.median(heights) if heights else 12
        prev_bottom = None
        for top, typ, p in items:
            if typ == "table":
                blocks.append(("table", p))
                prev_bottom = None
                continue
            text, bottom, ph, bold = p
            edge = top < ph * 0.1 or top > ph * 0.9
            if not text or (PAGE_NO_RE.match(text) and edge) or norm(text) in repeated or NOISE_RE.search(text):
                continue
            if typ in ("foot", "tabline"):
                blocks.append((typ, text))
                prev_bottom = bottom if typ == "tabline" else prev_bottom
                continue
            if prev_bottom is not None and top - prev_bottom > lh * 0.8:
                blocks.append(("gap", ""))
            blocks.append(("line", (text, bold)))
            prev_bottom = bottom
    return blocks


START_RE = re.compile(r"^(\d{1,2}\.?\s+)?(Grundsätze|Grundsatz|Einleitung|Anwendungsbereich|Geltungsbereich)$", re.I)
START_EN_RE = re.compile(r"^(\d{1,2}(\.\d{1,2})*\.?\s+)?(Principles?|Introduction|Introduction and purpose|"
                         r"Introduction \(background\)|Scope|Purpose|Objectives?|Background|Glossary)$", re.I)
TOC_RE = re.compile(r"^(inhaltsverzeichnis|table of contents|contents)$", re.I)


def strip_front_matter(blocks, start_re=START_RE):
    """Deckblatt, Status/Historie und Inhaltsverzeichnis weglassen.
    Der Richtlinientext beginnt mit 'Grundsätze', 'Einleitung', 'Anwendungsbereich' o.ä.;
    gibt es ein Inhaltsverzeichnis, zählt erst das zweite Vorkommen dieser Überschrift."""
    key = lambda b: re.sub(r"\s*\.{3,}.*$|\s+\d+$", "", b[1][0]).strip().lower() if b[0] == "line" else None
    toc = next((i for i, b in enumerate(blocks) if b[0] == "line" and TOC_RE.match(key(b) or "")), None)
    first = next((i for i, b in enumerate(blocks) if (toc is None or i > toc) and start_re.match(key(b) or "")), None)
    if first is None:
        return blocks
    if toc is not None:
        k = key(blocks[first])
        second = next((i for i in range(first + 1, len(blocks)) if key(blocks[i]) == k), None)
        first = second if second is not None else first
    return blocks[first:]


PARA_NO_RE = re.compile(r"^(\d{1,2}\.\d{1,3}\.?|\d{1,2}\.)\s+\S")
ICH_H2_RE = re.compile(r"^\d{1,2}\.\d(\.\d)?\s+[A-ZÄÖÜ(]")


EN_HEAD_RE = re.compile(r"^\d{1,2}(\.\d{1,2}){0,3}\.?\s+[A-Z][^.;:]{2,80}$")


def blocks_to_paragraphs(blocks, en: bool = False) -> list[tuple[str, str]]:
    """Zeilen zu Absätzen zusammenführen; Überschriften und Fußnoten erkennen.
    Rückgabe: (typ, text) mit typ in {'h1','h2','p','table_row'}."""
    lines = [b[1][0] for b in blocks if b[0] == "line"]
    # ICH-Schema (Teil II, Teil IV): Absätze 1.10, 1.11 …, Unterkapitel 1.1, 1.2 …
    ich = sum(bool(re.match(r"^\d{1,2}\.\d{2}\s+[A-ZÄÖÜ]", l)) for l in lines) >= 10
    paras, buf, feet = [], [], []
    last_h1, gap, heading_open = 0, False, False

    def flush():
        if buf:
            paras.append(("p", " ".join(buf).strip()))
            buf.clear()
        # Fußnoten hinter den Absatz stellen, in dem sie stehen, statt ihn zu zerreißen
        for f in feet:
            paras.append(("p", "Fußnote " + f))
        feet.clear()

    def add_heading(kind, text):
        nonlocal heading_open
        flush()
        paras.append((kind, text))
        heading_open = True

    for typ, payload in blocks:
        if typ == "gap":
            gap = True
            continue
        if typ == "foot":
            if re.match(r"^\d{1,3}\s", payload) or not feet:
                feet.append(payload)
            else:
                feet[-1] += " " + payload
            continue
        if typ == "tabline":
            if not (paras and paras[-1][0] == "table_line") or buf:
                flush()
                paras.append(("p", "Tabelle (Zeilen wie im Original, Spalten durch Leerzeichen getrennt):"))
            paras.append(("table_line", payload))
            heading_open = gap = False
            continue
        if typ == "table":
            flush()
            paras.extend(("table_row", s) for s in table_to_sentences(payload))
            heading_open = gap = False
            continue
        line, bold = payload
        # zweizeilige Überschrift fortsetzen ("… Zellkultu-" / "ren/Fermentation …")
        if heading_open and paras and paras[-1][0] in ("h1", "h2") and not PARA_NO_RE.match(line) and \
                (paras[-1][1].endswith("-") or line[:1].islower()) and len(line) < 90:
            t = paras[-1][1]
            paras[-1] = (paras[-1][0], (t[:-1] if t.endswith("-") and line[:1].islower() else t + " ") + line
                         if not t.endswith("-") or line[:1].islower() else t + line)
            gap = False
            continue
        heading_open = False
        if gap:
            gap = False
            # Lücke mitten im Satz (Seitenwechsel, Fußnote) ist kein Absatzende
            if not (buf and not re.search(r"[.:;!?]$", buf[-1]) and line[:1].islower()):
                flush()
        m1 = H1_RE.match(line)
        plain = not re.search(r"[.;,]$", line)
        # Kapitel = fortlaufend nummeriert und fett (bzw. im ICH-Schema kurz ohne Satzende)
        if m1 and int(m1.group(1)) == last_h1 + 1 and (bold or (len(line) < 60 and not buf)):
            last_h1 = int(m1.group(1))
            add_heading("h1", line)
            continue
        if bold and len(line) < 100 and re.match(r"^(Annex|ANNEX|PART|Part)\s+([IVX]+|[A-D]|\d+)\b", line):
            add_heading("h1", line)
            continue
        if ich and ICH_H2_RE.match(line) and len(line) < 100 and plain and \
                int(line.split(".")[0]) == last_h1:
            add_heading("h2", line)
            continue
        # Unterüberschrift = fett, kurz, ohne Satzzeichen am Ende ("Allgemeines", "Räumlichkeiten")
        if bold and len(line) < 100 and plain and not line.endswith(".") and line[:1].isupper() and \
                (not buf or re.search(r"[.:;]$", buf[-1])):
            add_heading("h2", line)
            continue
        if START_RE.match(line) and not buf:
            add_heading("h2", line)
            continue
        # englische Leitlinien: "1.2. Scope", "4.3 Documentation" (kurz, ohne Satzende) = Überschrift
        if en and EN_HEAD_RE.match(line) and len(line) < 80 and (not buf or re.search(r"[.:;]$", buf[-1])):
            add_heading("h2", line)
            continue
        if PARA_NO_RE.match(line) or re.match(r"^([•\-–▪]|\(?[a-z]{1,2}\)|\([ivx]{1,5}\)|[ivx]{1,5}\.)\s", line):
            flush()  # nummerierter Absatz / Aufzählungspunkt beginnt neuen Absatz
        if buf and buf[-1].endswith("-") and not re.match(r"^(und|oder|bzw\.|sowie)\b", line):
            # Silbentrennung auflösen ("Ausgangs-/material"); vor Großbuchstaben Bindestrich behalten ("API-/Ausgangsmaterial")
            buf[-1] = (buf[-1] if line[:1].isupper() else buf[-1][:-1]) + line
            continue
        buf.append(line)
    flush()
    # kurze, allein stehende Zeilen ohne Satzzeichen sind Zwischenüberschriften
    return [("h2", t) if typ == "p" and len(t) < 70 and len(t.split()) <= 8 and t[:1].isupper()
            and not re.search(r"[.:;,]$", t) and not t.startswith("Fußnote") else (typ, t)
            for typ, t in paras]


# --------------------------------------------------------------------------- Aufteilen
def words(paras) -> int:
    return sum(len(t.split()) for _, t in paras)


def split_sections(paras, level="h1"):
    sections, cur = [], []
    for p in paras:
        if p[0] == level and cur:
            sections.append(cur)
            cur = []
        cur.append(p)
    if cur:
        sections.append(cur)
    return sections


def chunk(paras):
    """Kapitel zusammenfassen, bis MAX_WORDS erreicht; zu große Kapitel an Unterkapiteln teilen."""
    if words(paras) <= MAX_WORDS:
        return [paras]
    units = []
    for sec in split_sections(paras, "h1"):
        if words(sec) <= MAX_WORDS:
            units.append(sec)
        else:
            subs = split_sections(sec, "h2")
            heading = sec[0] if sec[0][0] == "h1" else None
            for i, s in enumerate(subs):
                if i > 0 and heading:  # Kapitelüberschrift als Kontext wiederholen
                    s = [("h1", heading[1] + " (Fortsetzung)")] + s
                units.append(s)
    # Abschnitte ohne Unterüberschriften, die deutlich zu lang sind, absatzweise teilen
    split_units = []
    for u in units:
        if words(u) <= MAX_WORDS * 1.15:
            split_units.append(u)
            continue
        head = u[0] if u[0][0] in ("h1", "h2") else None
        cur_u = []
        for p_ in u:
            if cur_u and words(cur_u) + len(p_[1].split()) > MAX_WORDS and p_[0] not in ("h1", "h2"):
                split_units.append(cur_u)
                cur_u = [(head[0], head[1] + " (Fortsetzung)")] if head else []
            cur_u.append(p_)
        split_units.append(cur_u)
    units = split_units
    parts, cur = [], []
    for u in units:
        if cur and words(cur) + words(u) > MAX_WORDS:
            parts.append(cur)
            cur = []
        cur = cur + u
    if cur:
        parts.append(cur)
    # Kleinstteile (z. B. nur der Titel eines Rechtsakts) mit dem Nachbarteil zusammenlegen
    merged = []
    for part in parts:
        if merged and words(merged[-1]) < 400:
            merged[-1] = merged[-1] + part
        else:
            merged.append(part)
    if len(merged) > 1 and words(merged[-1]) < 400:
        merged[-2] = merged[-2] + merged.pop()
    return merged


def topic_of(part) -> str:
    heads = []
    for t, text in part:
        if t == "h1" and "(Fortsetzung)" not in text:
            h = re.sub(r"^(KAPITEL|TITEL|ANHANG|TEIL)\s+[IVXLC\d]+\s*", "", text, flags=re.I)
            h = re.sub(r"\s*\([^)]*\)", "", re.sub(r"^\d{1,2}\.?\s+", "", h))  # "(RBA)" u.ä. weglassen
            if not h:
                continue
            heads.append(h if len(h) <= 35 else trim_words(h[:35].rsplit(" ", 1)[0]))
    if not heads:
        heads = [re.sub(r"^\d{1,2}\.\d{1,2}\.?\s+", "", t2) for t, t2 in part if t == "h2"][:2]
    nums = [int(H1_RE.match(t2).group(1)) for t, t2 in part if t == "h1" and H1_RE.match(t2)]
    rng = ""
    labels = [re.match(r"^(KAPITEL|ANHANG)\s+([IVXLC]+)\b", t2, re.I) for t, t2 in part if t == "h1"]
    labels = [m for m in labels if m]
    if labels and not nums:
        kind = "Kap" if labels[0].group(1).upper() == "KAPITEL" else "Anhang"
        same = [m for m in labels if (m.group(1).upper() == "KAPITEL") == (kind == "Kap")]
        a, b = same[0].group(2), same[-1].group(2)
        rng = f"{kind} {a}" if a == b or roman(b) < roman(a) else f"{kind} {a}-{b}"
    if nums:
        rng = f"Kap {nums[0]}" if nums[0] == nums[-1] else f"Kap {nums[0]}-{nums[-1]}"
    topic = ", ".join(heads[:3])  # der Kapitelbereich zeigt, dass ggf. mehr enthalten ist
    return f"{rng} {topic}".strip()


STOP = {"und", "oder", "der", "die", "das", "des", "von", "für", "mit", "zur", "zum", "bei", "in", "im", "zu"}


def roman(s: str) -> int:
    vals, total, prev = {"I": 1, "V": 5, "X": 10, "L": 50, "C": 100}, 0, 0
    for ch in reversed(s.upper()):
        v = vals.get(ch, 0)
        total, prev = (total - v, prev) if v < prev else (total + v, v)
    return total


def trim_words(s: str) -> str:
    words = s.rstrip(" ,—–-:").split()
    while len(words) > 1 and words[-1].lower() in STOP:
        words.pop()
    return " ".join(words).rstrip(" ,—–-:")


def safe_name(s: str, limit: int = 110) -> str:
    s = re.sub(r'[<>:"/\\|?*]', "-", s)
    s = re.sub(r"\s+", " ", s).strip()
    if len(s) > limit:  # an Wortgrenze kürzen
        s = trim_words(s[:limit].rsplit(" ", 1)[0])
    return s


# --------------------------------------------------------------------------- Word
def write_docx(path: Path, title: str, source: str, part):
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name, st.font.size = "Calibri", Pt(11)
    doc.core_properties.title = title
    doc.core_properties.subject = source
    doc.add_heading(title, level=0)
    doc.add_paragraph(f"Quelle: {source}").italic = True
    in_table = False
    for typ, text in part:
        if typ == "h1":
            doc.add_heading(text, level=1)
        elif typ == "h2":
            doc.add_heading(text, level=2)
        elif typ == "table_line":
            doc.add_paragraph(text)
        elif typ == "table_row":
            if not in_table:
                doc.add_paragraph("Tabelle (zeilenweise wiedergegeben):").runs[0].bold = True
            doc.add_paragraph(text, style="List Bullet")
        else:
            doc.add_paragraph(text)
        in_table = typ == "table_row"
    doc.save(path)


def process(key: str, pdf: Path, out: Path):
    base = TITLES.get(key, key)
    paras = blocks_to_paragraphs(strip_front_matter(extract_blocks(pdf)))
    parts = chunk(paras)
    source = f"EU-GMP-Leitfaden, {key} (Bekanntmachung des BMG), {SOURCE_PAGE}"
    files = []
    for i, part in enumerate(parts, 1):
        if len(parts) == 1:
            name = base
        else:
            name = f"{base} ({i} von {len(parts)}) {topic_of(part)}"
        name = safe_name(name)
        write_docx(out / f"{name}.docx", name, source, part)
        files.append((name, words(part)))
    return files


# --------------------------------------------------------------------------- EUR-Lex (HTML/XHTML)
def _text(el, refs=None) -> str:
    """Text eines Elements; Fußnotenverweise "( 1 )" im Fließtext werden entfernt und in refs gesammelt."""
    is_note = re.search(r"(^|\s)(oj-)?note$", _cls(el)) if el.name == "p" else False
    if not is_note:
        for tag in el.select("span.oj-note-tag, a.oj-note-tag, span.note-tag"):
            if refs is not None:
                refs.extend(a.get("href", "").lstrip("#") for a in tag.find_all("a") if a.get("href"))
            tag.decompose()
    t = " ".join(el.get_text(" ", strip=True).split())
    t = t.replace(" ,", ",").replace(" .", ".").replace(" ;", ";").replace("( ", "(").replace(" )", ")")
    return re.sub(r"\s*\(\s*\)", "", t)


def _cls(el) -> str:
    return " ".join(el.get("class") or [])


def eurlex_to_paragraphs(html: str) -> list[tuple[str, str]]:
    """Amtsblatt-HTML (EUR-Lex/Cellar) -> Absätze. Kapitel/Anhänge = h1, Abschnitte/Artikel = h2,
    Aufzählungen ('(a) …') als eigene Absätze, echte Tabellen zeilenweise, Fußnoten ans Ende."""
    from bs4 import BeautifulSoup

    soup = BeautifulSoup(html, "html.parser")
    root = soup.body or soup
    paras, notes, refs_of = [], {}, {}  # notes: Anker-ID -> Text; refs_of: Absatzindex -> [Anker-IDs]
    pending = {"sec": None, "art": None}
    seen_title = [False]

    def emit(kind, text):
        if text:
            paras.append((kind, text))

    def walk(node, prefix=""):
        for el in node.children:
            if not getattr(el, "name", None):
                continue
            c = _cls(el)
            if el.name == "table":
                table(el, prefix)
                prefix = ""
                continue
            if el.name != "p":
                walk(el, prefix)
                prefix = ""
                continue
            refs = []
            t = _text(el, refs)
            if not t or re.search(r"(^|\s)(oj-)?hd-", " " + c) or "signatory" in c:
                continue
            if refs:
                refs_of.setdefault(len(paras), []).extend(refs)
            if re.search(r"(^|\s)(oj-)?note$", c):
                ids = [a.get("id") for a in el.find_all("a") if a.get("id")]
                notes[ids[0] if ids else f"n{len(notes)}"] = t
                continue
            elif re.search(r"(oj-)?doc-ti", c):
                if not seen_title[0]:  # Titel des Rechtsakts: ein Absatz
                    if paras and paras[-1][0] == "title":
                        paras[-1] = ("title", paras[-1][1] + " " + t)
                    else:
                        paras.append(("title", t))
                elif re.match(r"^(ANHANG|ANLAGE|TEIL)\b", t, re.I) or not paras or paras[-1][0] != "h1":
                    emit("h1", t)
                else:
                    paras[-1] = ("h1", paras[-1][1] + " " + t)
            elif re.search(r"ti-section-1", c):
                pending["sec"] = t
            elif re.search(r"ti-section-2", c) and pending["sec"]:
                full = f"{pending['sec']} {t}"
                emit("h1" if re.match(r"^(KAPITEL|TITEL|TEIL)\b", pending["sec"], re.I) else "h2", full)
                pending["sec"] = None
            elif re.search(r"(^|\s)(oj-)?ti-art", c):
                pending["art"] = t
                emit("h2", t)
            elif re.search(r"(^|\s)(oj-)?sti-art", c) and paras and paras[-1] == ("h2", pending["art"]):
                paras[-1] = ("h2", f"{pending['art']} {t}")
            elif re.search(r"ti-grseq|ti-tbl", c):
                emit("h2", t)
            else:
                seen_title[0] = True
                if t.lower().startswith("in erwägung nachstehender gründe"):
                    emit("h1", "Erwägungsgründe")
                emit("p", (prefix + " " + t).strip() if prefix else t)
                if re.search(r"(HAT|HABEN) FOLGENDE .* ERLASSEN:?$", t):
                    emit("h1", "Verfügender Teil")
                prefix = ""
            if not re.search(r"(oj-)?doc-ti", c):
                seen_title[0] = True

    def table(tbl, prefix=""):
        rows = [tr for tr in tbl.find_all("tr") if tr.find_parent("table") is tbl]
        cells = [[td for td in tr.find_all(["td", "th"]) if td.find_parent("tr") is tr] for tr in rows]
        # Aufzählungstabelle: 2 Spalten, erste Spalte nur Marke wie "(a)", "1.", "—"
        if cells and all(len(r) == 2 and len(_text(r[0])) <= 8 for r in cells if r):
            for r in cells:
                if r:
                    mark = _text(r[0])
                    walk(r[1], mark)
            return
        data = [[_text(td) for td in r] for r in cells if r]
        if any("Amtsblatt der Europäischen Union" in c for r in data[:1] for c in r):
            return  # Kopfzeile des Amtsblatts
        if prefix:
            emit("p", prefix)
        paras.extend(("table_row", x) for x in table_to_sentences(data))

    walk(root)
    # Fußnoten direkt hinter den Absatz stellen, der auf sie verweist
    out, used = [], set()
    for i, (k, t) in enumerate(paras):
        # leere Gliederungsebene entfernen (z. B. "Verfügender Teil" direkt vor "KAPITEL I")
        if k == "h1" and i + 1 < len(paras) and paras[i + 1][0] == "h1" and t == "Verfügender Teil":
            continue
        out.append(("p", t) if k == "title" else (k, t))
        for r in refs_of.get(i, []):
            if r in notes and r not in used:
                used.add(r)
                out.append(("p", "Fußnote " + notes[r]))
    rest = [n for key, n in notes.items() if key not in used]
    if rest:  # nicht zuordenbare Fußnoten ans Ende
        out.append(("h2", "Fußnoten"))
        out.extend(("p", "Fußnote " + n) for n in rest)
    return out


# --------------------------------------------------------------------------- Word-Vorlagen (.docx)
def docx_to_paragraphs(path: Path) -> list[tuple[str, str]]:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    d = Document(path)
    paras = []
    for el in d.element.body.iterchildren():
        tag = el.tag.split("}")[-1]
        if tag == "p":
            p = Paragraph(el, d)
            t = " ".join(p.text.split())
            if not t:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""
            paras.append(("h1" if style.startswith(("heading 1", "title")) else
                          "h2" if style.startswith("heading") else "p", t))
        elif tag == "tbl":
            tbl = Table(el, d)
            rows = []
            for r in tbl.rows:
                row = []
                for c in r.cells:
                    txt = " ".join(c.text.split())
                    row.append(txt if not row or txt != row[-1] else "")  # verbundene Zellen nicht doppeln
                rows.append(row)
            paras.extend(("table_row", x) for x in table_to_sentences(rows))
    return paras


# --------------------------------------------------------------------------- Zusatzquellen (quellen.json)
def process_source(q: dict, path: Path, out: Path):
    base = q["titel"] + (" (EN)" if q.get("sprache") == "EN" else "")
    if path.suffix == ".pdf":
        start = re.compile(q["start"], re.I) if q.get("start") else (START_EN_RE if q.get("sprache") == "EN" else START_RE)
        blocks = extract_blocks(path, line_numbers=q.get("zeilennummern", False))
        if q.get("start") != "":
            blocks = strip_front_matter(blocks, start)
        paras = blocks_to_paragraphs(blocks, en=q.get("sprache") == "EN")
    elif path.suffix == ".docx":
        paras = docx_to_paragraphs(path)
    else:
        paras = eurlex_to_paragraphs(path.read_text(encoding="utf-8", errors="replace"))
    if q["typ"] == "eurlex":
        source = f"Amtsblatt der EU, CELEX {q['celex']}, https://eur-lex.europa.eu/legal-content/DE/TXT/?uri=CELEX:{q['celex']}"
    else:
        source = q["url"]
    parts = chunk(paras)
    files = []
    for i, part in enumerate(parts, 1):
        name = safe_name(base if len(parts) == 1 else f"{base} ({i} von {len(parts)}) {topic_of(part)}")
        write_docx(out / f"{name}.docx", name, source, part)
        files.append((name, words(part)))
    return files


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--out", default="ausgabe")
    ap.add_argument("--pdf-dir", help="Ordner mit bereits geladenen PDFs (Dateiname = Linktext, z.B. 'Kapitel 1.pdf')")
    ap.add_argument("--quellen", help="quellen.json mit Zusatzdokumenten (vorher mit quellen_laden.py laden)")
    args = ap.parse_args()
    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)

    if args.pdf_dir == "-":
        pdfs = []  # nur Zusatzquellen verarbeiten
    elif args.pdf_dir:
        pdfs = [(p.stem, p) for p in sorted(Path(args.pdf_dir).glob("*.pdf"))]
    else:
        pdfs = download_pdfs(out / "_pdf")

    for key, pdf in pdfs:
        for name, w in process(key, pdf, out):
            print(f"  {w:6d} Wörter  {name}.docx")

    if args.quellen:
        import json
        src_dir = out / "_quellen"
        for q in json.loads(Path(args.quellen).read_text(encoding="utf-8"))["quellen"]:
            files = sorted(src_dir.glob(f"{q['key']}.*"))
            files = [f for f in files if f.suffix in (".pdf", ".html", ".docx")]
            if not files:
                print(f"  FEHLT: {q['key']} (nicht geladen, siehe _quellen/download.log)")
                continue
            for name, w in process_source(q, files[0], out):
                print(f"  {w:6d} Wörter  {name}.docx")

    with zipfile.ZipFile(out.with_suffix(".zip"), "w", zipfile.ZIP_DEFLATED) as z:
        for f in sorted(out.glob("*.docx")):
            z.write(f, f.name)
    print(f"Fertig: {out.with_suffix('.zip')}")


if __name__ == "__main__":
    main()
